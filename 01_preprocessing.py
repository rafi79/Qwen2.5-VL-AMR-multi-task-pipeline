"""
╔══════════════════════════════════════════════════════════════════════════╗
║  Qwen2.5-VL AMR — PART 1: Data Preprocessing & Cleaning Pipeline        ║
║  Dataset: https://datadryad.org/dataset/doi:10.5061/dryad.5dv41nsfj     ║
╠══════════════════════════════════════════════════════════════════════════╣
║  Install dependencies before running:                                    ║
║    !pip install einops timm -q                                           ║
║    !pip install python-docx openpyxl -q                                  ║
║    !pip install bitsandbytes accelerate -q                               ║
║    !pip install peft einops timm -q                                      ║
║    !pip install peft --quiet                                             ║
╠══════════════════════════════════════════════════════════════════════════╣
║  Kaggle dataset paths:                                                   ║
║    /kaggle/input/datasets/abedhossain/dryad-image-dataset-of-amr/        ║
║      ├── images_original/1.1.1. original.jpg                            ║
║      ├── images_measured/1.1.1. measured.jpg                            ║
║      └── Tables/                                                         ║
║          ├── Overview GPT JCM.xlsx   ← master mechanism labels          ║
║          └── Table 1.1.1..docx       ← per-sample antibiotic+zone+S/I/R ║
╠══════════════════════════════════════════════════════════════════════════╣
║  Steps:                                                                  ║
║    1. Parse Overview xlsx  → mechanism labels per sample                 ║
║    2. Parse all .docx files → per-antibiotic zone + S/I/R labels         ║
║    3. Merge both into one clean master DataFrame                         ║
║    4. Verify image files exist for every sample                          ║
║    5. Report cleaning decisions and class distribution                   ║
║    6. Save: master_labels.csv + image_index.csv                          ║
╠══════════════════════════════════════════════════════════════════════════╣
║  Outputs → data/processed/                                               ║
║    master_labels.csv     (one row per antibiotic record)                 ║
║    image_index.csv       (one row per image / sample)                    ║
║    class_distribution.png                                                ║
╚══════════════════════════════════════════════════════════════════════════╝
"""

import os
import re
import glob
from pathlib import Path

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from docx import Document           # pip install python-docx
from openpyxl import load_workbook  # pip install openpyxl

# ── Paths ──────────────────────────────────────────────────────────────────
BASE       = Path("/kaggle/input/datasets/abedhossain/dryad-image-dataset-of-amr")
DIR_ORIG   = BASE / "images_original"
DIR_MEAS   = BASE / "images_measured"
DIR_TABLES = BASE / "Tables"
OVERVIEW   = DIR_TABLES / "Overview GPT JCM.xlsx"

OUT_DIR = Path("data/processed")
OUT_DIR.mkdir(parents=True, exist_ok=True)


# ══════════════════════════════════════════════════════════════════════════
# STEP 1 — Parse Overview xlsx (mechanism labels)
# ══════════════════════════════════════════════════════════════════════════
def parse_overview_xlsx(path: Path) -> pd.DataFrame:
    """
    Reads the master Excel file containing one row per sample with columns:
      IMAGE | SPECIES | NONE | ESBL | AMPC | CARBAPENEMASE
    Derives the multi-class ResistanceMechanism label from the boolean flags.
    """
    df = pd.read_excel(path)

    # Print raw column names so we can verify the parse
    print(f"Raw columns: {list(df.columns)}")

    # Rename columns explicitly
    df.columns = ["IMAGE", "SPECIES", "NONE", "ESBL", "AMPC", "CARBAPENEMASE"]
    print(f"Overview rows: {len(df)}")

    # Clean Image ID
    df["IMAGE"] = df["IMAGE"].astype(str).str.strip().str.rstrip(".")

    # Clean SPECIES
    df["SPECIES"] = df["SPECIES"].astype(str).str.strip()

    # Boolean map — handles all variants seen in the xlsx
    BOOL_MAP = {
        "yes": True,  "no": False,
        "not applicable": False,  "not tested": np.nan,
        "not assessable": np.nan, "?": np.nan,
        "nan": np.nan, "": np.nan,
    }

    for col in ["NONE", "ESBL", "AMPC", "CARBAPENEMASE"]:
        df[col] = (df[col].astype(str).str.strip().str.lower().map(BOOL_MAP))

    def derive_mechanism(row):
        active = [
            m for m, v in [
                ("ESBL", row["ESBL"]),
                ("AmpC", row["AMPC"]),
                ("Carbapenemase", row["CARBAPENEMASE"]),
            ]
            if v is True
        ]
        if len(active) == 0:
            return "None" if row["NONE"] is True else "Unknown"
        elif len(active) == 1:
            return active[0]
        else:
            return "Combination"

    df["ResistanceMechanism"] = df.apply(derive_mechanism, axis=1)

    print("\nMechanism distribution:")
    print(df["ResistanceMechanism"].value_counts())

    return df[["IMAGE", "SPECIES", "NONE", "ESBL", "AMPC",
               "CARBAPENEMASE", "ResistanceMechanism"]]


# ══════════════════════════════════════════════════════════════════════════
# STEP 2 — Parse per-sample .docx files (antibiotic + zone + S/I/R)
# ══════════════════════════════════════════════════════════════════════════
def parse_one_docx(docx_path: Path) -> pd.DataFrame:
    """
    Each .docx (e.g. Table 1.1.1..docx) contains a table:
      Abbreviation | Antibiotic | Inhibition zone (mm) | Interpretation | ...

    Returns DataFrame:
      SampleID | Abbreviation | Antibiotic | ZoneDiameter_mm | Interpretation
    """
    stem      = docx_path.stem                     # "Table 1.1.1."
    sample_id = stem.replace("Table ", "").strip().rstrip(".")

    try:
        doc = Document(docx_path)
    except Exception as e:
        print(f"  Could not open {docx_path.name}: {e}")
        return pd.DataFrame()

    rows = []
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            cells = [c.text.strip() for c in row.cells]
            if i == 0:
                continue  # skip header
            if len(cells) < 4:
                continue

            abbrev     = cells[0] if len(cells) > 0 else ""
            antibiotic = cells[1] if len(cells) > 1 else ""
            zone_raw   = cells[2] if len(cells) > 2 else ""
            interp     = cells[3] if len(cells) > 3 else ""

            if not abbrev and not antibiotic:
                continue

            try:
                zone_mm = float(re.sub(r"[^\d.]", "", zone_raw))
            except (ValueError, TypeError):
                zone_mm = np.nan

            interp_clean = interp.strip().upper()
            if interp_clean not in ("S", "I", "R"):
                interp_clean = np.nan

            rows.append({
                "SampleID":        sample_id,
                "Abbreviation":    abbrev,
                "Antibiotic":      antibiotic,
                "ZoneDiameter_mm": zone_mm,
                "Interpretation":  interp_clean,
            })

    return pd.DataFrame(rows)


def parse_all_docx(tables_dir: Path) -> pd.DataFrame:
    """Parse all Table X.X.X..docx files in the Tables directory."""
    docx_files = sorted(tables_dir.glob("Table *.docx"))
    print(f"\nFound {len(docx_files)} .docx table files")

    all_dfs = []
    for i, f in enumerate(docx_files):
        df = parse_one_docx(f)
        if not df.empty:
            all_dfs.append(df)
        if (i + 1) % 20 == 0:
            print(f"  Parsed {i+1}/{len(docx_files)}...")

    if not all_dfs:
        print("WARNING: No docx data parsed.")
        return pd.DataFrame()

    combined = pd.concat(all_dfs, ignore_index=True)
    print(f"\nTotal antibiotic records parsed: {len(combined)}")
    print(f"Unique samples                 : {combined['SampleID'].nunique()}")
    print(f"\nInterpretation distribution:")
    print(combined["Interpretation"].value_counts())
    print(f"\nMissing zone diameters: {combined['ZoneDiameter_mm'].isna().sum()}")

    return combined


# ══════════════════════════════════════════════════════════════════════════
# STEP 3 — Build image index
# ══════════════════════════════════════════════════════════════════════════
def build_image_index(overview_df: pd.DataFrame) -> pd.DataFrame:
    """
    For each sample in overview, verify original + measured image paths.
    Filename format: "1.1.1. original.jpg"
    """
    records      = []
    missing_orig = []
    missing_meas = []

    for _, row in overview_df.iterrows():
        sid      = row["SampleID"]
        sid_dot  = sid + "."

        orig_path  = DIR_ORIG / f"{sid_dot} original.jpg"
        meas_path  = DIR_MEAS / f"{sid_dot} measured.jpg"
        orig_exist = orig_path.exists()
        meas_exist = meas_path.exists()

        if not orig_exist:
            missing_orig.append(sid)
        if not meas_exist:
            missing_meas.append(sid)

        records.append({
            "SampleID":        sid,
            "original_path":   str(orig_path) if orig_exist else None,
            "measured_path":   str(meas_path) if meas_exist else None,
            "original_exists": orig_exist,
            "measured_exists": meas_exist,
        })

    df = pd.DataFrame(records)
    print(f"\nImage index: {len(df)} samples")
    print(f"  Original images found: {df['original_exists'].sum()}")
    print(f"  Measured images found: {df['measured_exists'].sum()}")
    if missing_orig:
        short = missing_orig[:5]
        print(f"  Missing originals    : {short}{'...' if len(missing_orig)>5 else ''}")

    return df


# ══════════════════════════════════════════════════════════════════════════
# STEP 4 — Merge everything into master DataFrame
# ══════════════════════════════════════════════════════════════════════════
def build_master_df(
    overview_df: pd.DataFrame,
    antibiotic_df: pd.DataFrame,
    image_index_df: pd.DataFrame,
) -> pd.DataFrame:
    master = overview_df.merge(image_index_df, on="SampleID", how="left")

    master_full = antibiotic_df.merge(
        master[[
            "SampleID", "SPECIES", "ResistanceMechanism",
            "original_path", "measured_path",
            "original_exists", "measured_exists",
        ]],
        on="SampleID",
        how="left",
    )

    print(f"\nMaster DataFrame shape: {master_full.shape}")
    return master_full


# ══════════════════════════════════════════════════════════════════════════
# STEP 5 — Final cleaning passes
# ══════════════════════════════════════════════════════════════════════════
def final_cleaning(df: pd.DataFrame) -> pd.DataFrame:
    original_len = len(df)
    report       = []

    # 5a. Drop rows with missing images
    before = len(df)
    df     = df[df["original_exists"] == True].copy()
    report.append(f"Dropped {before - len(df)} rows: no original image found")

    # 5b. Drop rows where Interpretation is NaN
    before = len(df)
    df     = df.dropna(subset=["Interpretation"])
    report.append(f"Dropped {before - len(df)} rows: missing S/I/R interpretation")

    # 5c. Drop rows where ZoneDiameter_mm is NaN or implausible
    before = len(df)
    df     = df.dropna(subset=["ZoneDiameter_mm"])
    df     = df[(df["ZoneDiameter_mm"] >= 6) & (df["ZoneDiameter_mm"] <= 50)]
    report.append(f"Dropped {before - len(df)} rows: missing/implausible zone diameter")

    # 5d. Standardise Abbreviation
    df["Abbreviation"] = df["Abbreviation"].str.upper().str.strip()

    # 5e. Standardise Antibiotic names
    df["Antibiotic"] = df["Antibiotic"].str.strip()

    # 5f. Flag Unknown mechanism
    unknown_count = (df["ResistanceMechanism"] == "Unknown").sum()
    report.append(f"Flagged {unknown_count} rows: 'Unknown' mechanism (kept)")

    # 5g. Numeric label
    df["label_int"] = df["Interpretation"].map({"S": 0, "I": 1, "R": 2})

    # 5h. Binary resistance
    df["is_resistant"] = (df["Interpretation"] == "R").astype(int)

    print("\n── CLEANING REPORT ──────────────────────────────────")
    for r in report:
        print(f"  • {r}")
    print(f"  Final rows: {len(df)} (from {original_len})")
    print("─────────────────────────────────────────────────────")

    return df


# ══════════════════════════════════════════════════════════════════════════
# STEP 6 — Visualise class balance
# ══════════════════════════════════════════════════════════════════════════
def visualise_distribution(df: pd.DataFrame):
    fig, axes = plt.subplots(1, 3, figsize=(15, 4))

    # S/I/R distribution
    ax = axes[0]
    counts = df["Interpretation"].value_counts()
    bars   = ax.bar(counts.index, counts.values,
                    color=["#1D9E75", "#EF9F27", "#E24B4A"])
    ax.set_title("S / I / R Distribution")
    ax.set_xlabel("Interpretation"); ax.set_ylabel("Count")
    for bar, count in zip(bars, counts.values):
        ax.text(bar.get_x() + bar.get_width()/2,
                bar.get_height() + 5, str(count),
                ha="center", va="bottom", fontsize=11)

    # Resistance mechanism
    ax = axes[1]
    mech_counts = (df.drop_duplicates("SampleID")["ResistanceMechanism"]
                   .value_counts())
    colors = {
        "None": "#1D9E75", "ESBL": "#EF9F27", "AmpC": "#378ADD",
        "Carbapenemase": "#E24B4A", "Combination": "#7F77DD",
        "Unknown": "#888780",
    }
    ax.bar(mech_counts.index, mech_counts.values,
           color=[colors.get(m, "#888780") for m in mech_counts.index])
    ax.set_title("Resistance Mechanism\n(per sample)")
    ax.set_xlabel("Mechanism"); ax.set_ylabel("Samples")
    plt.setp(ax.get_xticklabels(), rotation=20, ha="right")

    # Zone diameter distribution by S/I/R
    ax = axes[2]
    for interp, color in [("S", "#1D9E75"), ("I", "#EF9F27"), ("R", "#E24B4A")]:
        subset = df[df["Interpretation"] == interp]["ZoneDiameter_mm"]
        ax.hist(subset, bins=20, alpha=0.6, color=color, label=interp)
    ax.axvline(22, color="gray", linestyle="--", alpha=0.5, label="~breakpoint")
    ax.set_title("Zone Diameter Distribution\nby Interpretation")
    ax.set_xlabel("Zone Diameter (mm)"); ax.set_ylabel("Count")
    ax.legend()

    plt.tight_layout()
    plt.savefig(OUT_DIR / "class_distribution.png", dpi=150)
    print(f"\nSaved: {OUT_DIR}/class_distribution.png")
    plt.show()


# ══════════════════════════════════════════════════════════════════════════
# STEP 7 — Per-image summary (what VLM will receive)
# ══════════════════════════════════════════════════════════════════════════
def build_image_level_df(master_df: pd.DataFrame) -> pd.DataFrame:
    """One row per image, with summary stats + antibiotic context as JSON."""
    records = []
    for sid, grp in master_df.groupby("SampleID"):
        ab_context = grp[[
            "Abbreviation", "Antibiotic", "ZoneDiameter_mm", "Interpretation",
        ]].to_dict("records")

        records.append({
            "SampleID":            sid,
            "SPECIES":             grp["SPECIES"].iloc[0],
            "ResistanceMechanism": grp["ResistanceMechanism"].iloc[0],
            "original_path":       grp["original_path"].iloc[0],
            "measured_path":       grp["measured_path"].iloc[0],
            "n_antibiotics":       len(grp),
            "n_resistant":         (grp["Interpretation"] == "R").sum(),
            "n_susceptible":       (grp["Interpretation"] == "S").sum(),
            "n_intermediate":      (grp["Interpretation"] == "I").sum(),
            "antibiotic_context":  str(ab_context),
        })

    return pd.DataFrame(records)


# ══════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════
def run_preprocessing():
    print("=" * 60)
    print("  AMR VLM — DATA PREPROCESSING PIPELINE")
    print("=" * 60)

    print("\n[1/6] Parsing Overview xlsx...")
    overview_df = parse_overview_xlsx(OVERVIEW)
    overview_df = overview_df.rename(columns={"IMAGE": "SampleID"})

    print("\n[2/6] Parsing per-sample .docx files...")
    antibiotic_df = parse_all_docx(DIR_TABLES)

    print("\n[3/6] Building image index...")
    image_index_df = build_image_index(overview_df)

    print("\n[4/6] Merging datasets...")
    master_df = build_master_df(overview_df, antibiotic_df, image_index_df)

    print("\n[5/6] Final cleaning...")
    master_df = final_cleaning(master_df)

    print("\n[6/6] Visualising distributions...")
    visualise_distribution(master_df)

    master_df.to_csv(OUT_DIR / "master_labels.csv", index=False)
    print(f"\nSaved: {OUT_DIR}/master_labels.csv")

    image_df = build_image_level_df(master_df)
    image_df.to_csv(OUT_DIR / "image_index.csv", index=False)
    print(f"Saved: {OUT_DIR}/image_index.csv")

    print("\n" + "=" * 60)
    print("  FINAL DATASET SUMMARY")
    print("=" * 60)
    print(f"  Total images ready for VLM : {image_df['original_path'].notna().sum()}")
    print(f"  Total antibiotic records   : {len(master_df)}")
    print(f"  Unique antibiotics         : {master_df['Abbreviation'].nunique()}")
    print(f"  Species                    : {master_df['SPECIES'].nunique()}")
    print(f"\n  S/I/R breakdown:")
    for label, count in master_df["Interpretation"].value_counts().items():
        pct = count / len(master_df) * 100
        print(f"    {label}: {count} ({pct:.1f}%)")
    print(f"\n  Mechanism breakdown (images):")
    for mech, count in image_df["ResistanceMechanism"].value_counts().items():
        print(f"    {mech}: {count}")
    print("=" * 60)

    return master_df, image_df


if __name__ == "__main__":
    master_df, image_df = run_preprocessing()
